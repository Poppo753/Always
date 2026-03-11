from __future__ import annotations

import logging
import random
from pathlib import Path
from typing import Any

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN
from pptx.util import Emu, Pt

from yearbook.contracts.reasoning import DailySummary
from yearbook.pipeline.state import PipelineState

logger = logging.getLogger(__name__)

_MONTHS_EN = [
    '', 'January', 'February', 'March', 'April', 'May', 'June',
    'July', 'August', 'September', 'October', 'November', 'December',
]


def _hex(h: str) -> RGBColor:
    h = h.lstrip('#')
    return RGBColor(int(h[0:2], 16), int(h[2:4], 16), int(h[4:6], 16))


def _emu(total: int, frac: float) -> int:
    return int(total * frac)


def _parse_quote(raw: str) -> tuple[str, str]:
    """Split 'Ana: text' into ('Ana', 'text'). Returns ('', raw) if no prefix."""
    import re
    # Try common patterns: "Ana: text", "Botta: text", "Ana : text"
    m = re.match(r'^(Ana|Botta|ana|botta)\s*:\s*(.+)', raw, re.IGNORECASE)
    if m:
        speaker = m.group(1).capitalize()
        text = m.group(2).strip().strip('"\'')
        return speaker, text
    # No speaker found — return raw text with empty speaker
    return '', raw.strip().strip('"\'')


def _format_date_en(date_str: str) -> tuple[str, str]:
    """Return (day_number, month_name) from '2025-08-07'."""
    try:
        parts = date_str.split('-')
        day = str(int(parts[2]))
        month = _MONTHS_EN[int(parts[1])]
        return day, month
    except (IndexError, ValueError):
        return date_str, ''


class SlideConfig:
    def __init__(self, cfg: dict[str, Any]) -> None:
        slide_cfg = cfg.get('slide', {})
        self.W = slide_cfg.get('width_emu', 9144000)
        self.H = slide_cfg.get('height_emu', 5143500)

        colors = slide_cfg.get('colors', {})
        self.bg = _hex(colors.get('background', '1C1C2E'))
        self.title_color = _hex(colors.get('title', 'FFFFFF'))
        self.body_color = _hex(colors.get('body_text', 'E0E0E0'))
        self.accent = _hex(colors.get('accent', 'C9A84C'))
        self.quote_color = _hex(colors.get('quote_text', 'A8D8EA'))
        self.stats_color = _hex(colors.get('stats_text', '888888'))

        fonts = slide_cfg.get('fonts', {})
        self.title_pt = Pt(fonts.get('title_size', 28))
        self.body_pt = Pt(fonts.get('body_size', 13))
        self.quote_pt = Pt(fonts.get('quote_size', 14))
        self.stats_pt = Pt(fonts.get('stats_size', 10))

        layout = slide_cfg.get('layout', {})
        self.lay = layout

        collage_cfg = cfg.get('collage', {})
        self.max_images = collage_cfg.get('max_images', 4)


class PptxRenderer:
    def __init__(self, state: PipelineState) -> None:
        self.paths = state.paths
        self.scfg = SlideConfig(state.config.get('rendering', state.config))
        self._video_thumbs_by_date: dict[str, list[Path]] = self._build_video_thumb_index()
        self._day_images_by_date: dict[str, list[Path]] = self._build_day_image_index()
        self._emoji_by_date: dict[str, list[str]] = self._build_emoji_index()

    # ------------------------------------------------------------------
    # Index builders
    # ------------------------------------------------------------------

    def _build_video_thumb_index(self) -> dict[str, list[Path]]:
        import json
        from yearbook.contracts.media import VideoAnalysisRecord
        from yearbook.storage.jsonl_store import JsonlStore

        result: dict[str, list[Path]] = {}
        if not self.paths.media_registry.exists() or not self.paths.video_analysis.exists():
            return result

        reg = json.loads(self.paths.media_registry.read_text('utf-8'))
        date_by_id = {k: v.get('date', '') for k, v in reg.items()}

        thumbs_dir = self.paths.cache_video_frames
        for vid in JsonlStore(self.paths.video_analysis).read_all(VideoAnalysisRecord):
            date = date_by_id.get(vid.media_id, '')
            if not date:
                continue
            thumb = thumbs_dir / f'{vid.media_id}_thumb.jpg'
            if thumb.exists():
                result.setdefault(date, []).append(thumb)
        return result

    def _build_day_image_index(self) -> dict[str, list[Path]]:
        import json
        from yearbook.storage.jsonl_store import JsonlStore
        from yearbook.contracts.media import ImageAnalysisRecord

        result: dict[str, list[Path]] = {}
        if not self.paths.media_registry.exists():
            return result

        reg = json.loads(self.paths.media_registry.read_text('utf-8'))

        score_by_media: dict[str, float] = {}
        if self.paths.image_analysis.exists():
            for rec in JsonlStore(self.paths.image_analysis).read_all(ImageAnalysisRecord):
                score_by_media[rec.media_id] = rec.aesthetic_score

        bundles_dir = self.paths.run_dir / 'day_bundles'
        if not bundles_dir.exists():
            return result

        for bp in sorted(bundles_dir.glob('*.json')):
            bundle = json.loads(bp.read_text('utf-8'))
            date = bp.stem
            img_ids = bundle.get('image_analysis_ids', [])
            entries: list[tuple[float, Path]] = []
            for aid in img_ids:
                media_id = aid.replace('img_analysis_', '')
                rec = reg.get(media_id, {})
                rel = rec.get('relative_path')
                if rel:
                    p = self.paths.root / rel
                    if p.exists():
                        score = score_by_media.get(media_id, 0.0)
                        entries.append((score, p))
            entries.sort(key=lambda x: x[0], reverse=True)
            if entries:
                result[date] = [p for _, p in entries]
        return result

    def _build_emoji_index(self) -> dict[str, list[str]]:
        """Extract unique emoji per day from daily context messages."""
        import json, re
        result: dict[str, list[str]] = {}
        ctx_dir = self.paths.run_dir / 'daily_contexts'
        if not ctx_dir.exists():
            return result

        emoji_re = re.compile(
            r'[\U0001F600-\U0001F64F'
            r'\U0001F300-\U0001F5FF'
            r'\U0001F680-\U0001F6FF'
            r'\U0001F900-\U0001F9FF'
            r'\U0001FA00-\U0001FAFF'
            r'\U00002600-\U000026FF'
            r'\U00002700-\U000027BF'
            r'\U00002764'
            r']'
            r'[\U0001F3FB-\U0001F3FF]?'
            r'\uFE0F?',
            re.UNICODE,
        )

        for cp in sorted(ctx_dir.glob('*.json')):
            try:
                ctx = json.loads(cp.read_text('utf-8'))
                all_text = ' '.join(
                    m.get('text', '') for m in ctx.get('top_messages', [])
                )
                found = emoji_re.findall(all_text)
                seen: set[str] = set()
                unique: list[str] = []
                for e in found:
                    if e not in seen:
                        seen.add(e)
                        unique.append(e)
                if unique:
                    result[cp.stem] = unique
            except Exception:
                continue
        return result

    # ------------------------------------------------------------------
    # Quote validation
    # ------------------------------------------------------------------

    def _validate_quote(self, date: str, quote: str) -> str:
        """Return quote only if its text actually appears in that day's messages."""
        if not quote:
            return ''
        _, text = _parse_quote(quote)
        text_clean = text.lower().strip('"\'.,!? \u201c\u201d')
        if len(text_clean) < 4:
            return ''
        import json
        ctx_path = self.paths.run_dir / 'daily_contexts' / f'{date}.json'
        if not ctx_path.exists():
            return ''  # can't verify → drop
        try:
            ctx = json.loads(ctx_path.read_text('utf-8'))
        except Exception:
            return ''
        all_msg_text = ' '.join(
            m.get('text', '').lower() for m in ctx.get('top_messages', [])
        )
        # Check if at least 50% of meaningful words (3+ chars) appear in messages
        words = [w for w in text_clean.split() if len(w) >= 3]
        if not words:
            return quote  # very short, can't validate
        found = sum(1 for w in words if w in all_msg_text)
        if found / len(words) >= 0.5:
            return quote
        logger.debug('Dropping hallucinated quote for %s: %s', date, quote)
        return ''

    # ------------------------------------------------------------------
    # Main render
    # ------------------------------------------------------------------

    def render(self, summaries: list[DailySummary]) -> Presentation:
        prs = Presentation()
        prs.slide_width = Emu(self.scfg.W)
        prs.slide_height = Emu(self.scfg.H)
        blank = prs.slide_layouts[6]

        for idx, summary in enumerate(summaries):
            slide = prs.slides.add_slide(blank)
            self._set_background(slide)

            # Validate quote against actual messages
            summary.quote = self._validate_quote(summary.date, summary.quote)

            images = self._collect_images(summary)
            n = len(images)

            if n == 0:
                self._layout_text_only(slide, summary)
            elif n == 1:
                self._layout_hero(slide, summary, images[:1], idx)
            elif n == 2:
                self._layout_duo(slide, summary, images[:2], idx)
            else:
                self._layout_gallery(slide, summary, images[:self.scfg.max_images], idx)

            self._scatter_emoji(slide, summary.date, idx, has_images=(n > 0))

        return prs

    def _collect_images(self, summary: DailySummary) -> list[Path]:
        day_images = self._day_images_by_date.get(summary.date, [])
        video_thumbs = self._video_thumbs_by_date.get(summary.date, [])
        return (video_thumbs + day_images)[:self.scfg.max_images]

    # ------------------------------------------------------------------
    # Image helper -- preserves original quality & aspect ratio
    # ------------------------------------------------------------------

    def _add_picture_fit(self, slide, img_path: Path,
                         left: int, top: int, max_w: int, max_h: int):
        try:
            import io
            from PIL import Image as PILImage
            with PILImage.open(img_path) as img:
                iw, ih = img.size
                fmt = img.format
            if iw <= 0 or ih <= 0:
                return None
            ratio = iw / ih
            box_ratio = max_w / max_h if max_h > 0 else 1
            if ratio > box_ratio:
                w = max_w
                h = int(max_w / ratio)
            else:
                h = max_h
                w = int(max_h * ratio)
            x = left + (max_w - w) // 2
            y = top + (max_h - h) // 2
            # Convert unsupported formats (WebP) to PNG in memory
            if fmt and fmt.upper() not in ('BMP', 'GIF', 'JPEG', 'PNG', 'TIFF', 'WMF'):
                buf = io.BytesIO()
                with PILImage.open(img_path) as img:
                    img.save(buf, format='PNG')
                buf.seek(0)
                return slide.shapes.add_picture(buf, x, y, w, h)
            return slide.shapes.add_picture(str(img_path), x, y, w, h)
        except Exception as exc:
            logger.warning('Could not add image %s: %s', img_path, exc)
            return None

    # ------------------------------------------------------------------
    # Primitive element helpers
    # ------------------------------------------------------------------

    def _set_background(self, slide) -> None:
        fill = slide.background.fill
        fill.solid()
        fill.fore_color.rgb = self.scfg.bg

    def _add_date(self, slide, date_str: str,
                  left: int, top: int, width: int, align=None) -> None:
        """Styled date: '7  March' with accent line."""
        day_num, month = _format_date_en(date_str)
        W, H = self.scfg.W, self.scfg.H

        txb = slide.shapes.add_textbox(left, top, width, _emu(H, 0.07))
        tf = txb.text_frame
        tf.word_wrap = True
        p = tf.paragraphs[0]
        if align is not None:
            p.alignment = align

        r1 = p.add_run()
        r1.text = day_num
        r1.font.size = Pt(20)
        r1.font.bold = True
        r1.font.name = 'Segoe UI'
        r1.font.color.rgb = self.scfg.accent

        if month:
            r2 = p.add_run()
            r2.text = '  ' + month
            r2.font.size = Pt(12)
            r2.font.name = 'Segoe UI'
            r2.font.color.rgb = self.scfg.accent
            r2.font.italic = True

        # thin accent line below date
        line_y = top + _emu(H, 0.055)
        line_w = _emu(W, 0.06)
        if align == 1:
            line_x = left + (width - line_w) // 2
        elif align == 2:
            line_x = left + width - line_w
        else:
            line_x = left
        shape = slide.shapes.add_shape(
            MSO_SHAPE.RECTANGLE, line_x, line_y, line_w, Emu(18288)
        )
        shape.fill.solid()
        shape.fill.fore_color.rgb = self.scfg.accent
        shape.line.fill.background()

    def _add_title_at(self, slide, title: str,
                      left: int, top: int, width: int) -> None:
        txb = slide.shapes.add_textbox(left, top, width, _emu(self.scfg.H, 0.14))
        tf = txb.text_frame
        tf.word_wrap = True
        p = tf.paragraphs[0]
        run = p.add_run()
        run.text = title
        run.font.size = self.scfg.title_pt
        run.font.bold = True
        run.font.name = 'Segoe UI'
        run.font.color.rgb = self.scfg.title_color

    def _add_narrative(self, slide, summary: DailySummary,
                       left: int, top: int, width: int, height: int,
                       rng: random.Random | None = None,
                       font_size: Pt | None = None) -> int:
        """Add summary text only (no key_moments). Returns y after text."""
        size = font_size or self.scfg.body_pt
        txb = slide.shapes.add_textbox(left, top, width, height)
        tf = txb.text_frame
        tf.word_wrap = True
        p = tf.paragraphs[0]
        p.space_after = Pt(6)
        p.line_spacing = Pt(int(size) // 914 + 18)  # comfortable line spacing
        run = p.add_run()
        run.text = summary.summary
        run.font.size = size
        run.font.name = 'Segoe UI'
        run.font.color.rgb = self.scfg.body_color

        chars_per_line = max(20, int(width / (int(size) * 0.55 + 1)))
        summary_lines = max(2, len(summary.summary) // chars_per_line + 1)
        line_height_emu = int(size) + int(Pt(6))
        used_h = (summary_lines + 1) * line_height_emu
        return top + min(used_h, height)

    def _add_key_moments(self, slide, moments: list[str],
                         area_left: int, area_top: int,
                         area_width: int, area_height: int,
                         rng: random.Random) -> None:
        """Place each key moment as a separate textbox with organic positioning."""
        if not moments:
            return
        n = len(moments)
        row_h = area_height // max(n, 1)

        for i, moment in enumerate(moments):
            chars = len(moment)
            # Width proportional to text length, min 40%, max 95% of area
            need_frac = min(0.95, max(0.40, chars / 55))
            box_w = int(area_width * need_frac)

            # Horizontal: random within remaining space
            max_offset = max(0, area_width - box_w)
            x = area_left + rng.randint(0, max(1, max_offset))

            # Vertical: stacked with jitter
            base_y = area_top + row_h * i
            jitter_y = max(1, int(row_h * 0.12))
            y = base_y + rng.randint(-jitter_y, jitter_y)

            box_h = max(int(row_h * 0.80), int(Pt(14)))
            txb = slide.shapes.add_textbox(x, y, box_w, box_h)
            tf = txb.text_frame
            tf.word_wrap = True
            p = tf.paragraphs[0]
            run = p.add_run()
            run.text = '\u2022 ' + moment
            run.font.size = self._km_font_size
            run.font.name = 'Segoe UI'
            run.font.color.rgb = self.scfg.body_color

    def _add_quote(self, slide, quote: str,
                   left: int, top: int, width: int, align=None) -> None:
        if not quote:
            return
        speaker, text = _parse_quote(quote)

        txb = slide.shapes.add_textbox(left, top, width, _emu(self.scfg.H, 0.14))
        tf = txb.text_frame
        tf.word_wrap = True

        p = tf.paragraphs[0]
        if align is not None:
            p.alignment = align
        run = p.add_run()
        run.text = '\u201c' + text + '\u201d'
        run.font.size = self.scfg.quote_pt
        run.font.name = 'Segoe UI'
        run.font.italic = True
        run.font.color.rgb = self.scfg.quote_color

        if speaker:
            p2 = tf.add_paragraph()
            if align is not None:
                p2.alignment = align
            p2.space_before = Pt(2)
            run2 = p2.add_run()
            run2.text = speaker
            run2.font.size = Pt(11)
            run2.font.name = 'Segoe UI'
            run2.font.italic = True
            run2.font.color.rgb = self.scfg.accent

    def _scatter_emoji(self, slide, date: str, idx: int,
                       has_images: bool = False) -> None:
        """Place 1-5 emoji from that day's chat at random positions."""
        emojis = self._emoji_by_date.get(date, [])
        if not emojis:
            return

        W, H = self.scfg.W, self.scfg.H
        rng = random.Random(idx * 7919 + 17)

        if has_images:
            # Image layouts: right margin only (content on left + center)
            all_positions = [
                (0.93, 0.04), (0.88, 0.12), (0.92, 0.22),
                (0.87, 0.31), (0.94, 0.38), (0.89, 0.46),
                (0.93, 0.53), (0.87, 0.60), (0.92, 0.67),
                (0.88, 0.73), (0.94, 0.78), (0.90, 0.83),
                (0.86, 0.88), (0.93, 0.92), (0.88, 0.96),
            ]
        else:
            # Text-only: all four margins (content centered, margins free)
            all_positions = [
                # top edge
                (0.02, 0.02), (0.30, 0.01), (0.70, 0.01), (0.93, 0.02),
                # left edge
                (0.01, 0.20), (0.02, 0.45), (0.01, 0.70), (0.02, 0.88),
                # right edge
                (0.93, 0.18), (0.94, 0.42), (0.92, 0.65), (0.93, 0.85),
                # bottom edge
                (0.10, 0.93), (0.50, 0.94), (0.85, 0.92),
            ]

        pick_count = rng.randint(1, min(5, len(emojis)))
        picked_emojis = rng.sample(emojis, min(pick_count, len(emojis)))
        picked_positions = rng.sample(all_positions, pick_count)
        sizes = [13, 14, 15, 16, 18]

        for i, emoji_char in enumerate(picked_emojis):
            x_frac, y_frac = picked_positions[i]
            size = rng.choice(sizes)
            txb = slide.shapes.add_textbox(
                _emu(W, x_frac), _emu(H, y_frac),
                Emu(400000), Emu(350000),
            )
            tf = txb.text_frame
            p = tf.paragraphs[0]
            run = p.add_run()
            run.text = emoji_char
            run.font.size = Pt(size)

    # ------------------------------------------------------------------
    # Layout jitter helper
    # ------------------------------------------------------------------

    @staticmethod
    def _jf(base: float, jitter: float, rng: random.Random) -> float:
        """Jitter a fraction value by ±jitter, clamped to [0.01, 0.98]."""
        return max(0.01, min(0.98, base + rng.uniform(-jitter, jitter)))

    # ------------------------------------------------------------------
    # Layout: no media -- everything centred
    # ------------------------------------------------------------------

    def _layout_text_only(self, slide, summary: DailySummary) -> None:
        W, H = self.scfg.W, self.scfg.H
        rng = random.Random(hash(summary.date) * 3 + 100)
        self._km_font_size = Pt(12)  # bigger for text-only

        cx_f = self._jf(0.08, 0.02, rng)
        cw_f = self._jf(0.82, 0.03, rng)
        cx, cw = _emu(W, cx_f), _emu(W, cw_f)

        date_top = _emu(H, self._jf(0.04, 0.01, rng))
        title_top = _emu(H, self._jf(0.12, 0.01, rng))
        narr_top_f = self._jf(0.24, 0.02, rng)
        narr_top = _emu(H, narr_top_f)
        narr_h = _emu(H, 0.32)

        self._add_date(slide, summary.date, cx, date_top, cw, align=PP_ALIGN.CENTER)
        self._add_title_at(slide, summary.title, cx, title_top, cw)
        end_y = self._add_narrative(slide, summary, cx, narr_top, cw, narr_h, rng,
                                    font_size=Pt(15))

        # Key moments: full width, generous vertical space
        km_left = _emu(W, self._jf(0.06, 0.02, rng))
        km_width = _emu(W, self._jf(0.84, 0.03, rng))
        km_top = end_y + _emu(H, 0.03)
        km_bottom = _emu(H, 0.84) if not summary.quote else _emu(H, 0.76)
        km_h = km_bottom - km_top
        if km_h > 0:
            self._add_key_moments(slide, summary.key_moments, km_left, km_top, km_width, km_h, rng)

        self._add_quote(slide, summary.quote, cx, _emu(H, self._jf(0.86, 0.02, rng)), cw, align=PP_ALIGN.CENTER)

    # ------------------------------------------------------------------
    # Layout: hero -- one large image, alternating sides
    # ------------------------------------------------------------------

    def _layout_hero(self, slide, summary: DailySummary,
                     images: list[Path], idx: int) -> None:
        W, H = self.scfg.W, self.scfg.H
        rng = random.Random(hash(summary.date) * 3 + 201)
        self._km_font_size = Pt(11)
        right_img = (idx % 2 == 0)

        if right_img:
            tl_f = self._jf(0.04, 0.02, rng)
            tw_f = self._jf(0.46, 0.02, rng)
            il_f = self._jf(0.53, 0.02, rng)
            iw_f = self._jf(0.43, 0.02, rng)
        else:
            il_f = self._jf(0.04, 0.02, rng)
            iw_f = self._jf(0.43, 0.02, rng)
            tl_f = self._jf(0.50, 0.02, rng)
            tw_f = self._jf(0.46, 0.02, rng)

        tl, tw = _emu(W, tl_f), _emu(W, tw_f)
        il, iw = _emu(W, il_f), _emu(W, iw_f)

        date_top = _emu(H, self._jf(0.04, 0.01, rng))
        title_top = _emu(H, self._jf(0.12, 0.01, rng))
        img_top = _emu(H, self._jf(0.08, 0.02, rng))
        narr_top = _emu(H, self._jf(0.26, 0.02, rng))

        self._add_date(slide, summary.date, tl, date_top, tw)
        self._add_title_at(slide, summary.title, tl, title_top, tw)
        self._add_picture_fit(slide, images[0], il, img_top, iw, _emu(H, self._jf(0.78, 0.03, rng)))
        end_y = self._add_narrative(slide, summary, tl, narr_top, tw, _emu(H, 0.28), rng)

        km_top = end_y + _emu(H, 0.02)
        km_h = _emu(H, 0.70) - km_top + _emu(H, 0.04)
        if km_h > 0:
            self._add_key_moments(slide, summary.key_moments, tl, km_top, tw, km_h, rng)

        self._add_quote(slide, summary.quote, tl, _emu(H, self._jf(0.76, 0.02, rng)), tw)

    # ------------------------------------------------------------------
    # Layout: duo -- two images side by side
    # ------------------------------------------------------------------

    def _layout_duo(self, slide, summary: DailySummary,
                    images: list[Path], idx: int) -> None:
        W, H = self.scfg.W, self.scfg.H
        rng = random.Random(hash(summary.date) * 3 + 302)
        self._km_font_size = Pt(11)
        right_img = (idx % 2 == 0)
        gap = _emu(W, self._jf(0.015, 0.005, rng))

        if right_img:
            tl_f = self._jf(0.04, 0.02, rng)
            tw_f = self._jf(0.42, 0.02, rng)
            area_l_f = self._jf(0.50, 0.02, rng)
            each_w = (_emu(W, 0.46) - gap) // 2
        else:
            area_l_f = self._jf(0.04, 0.02, rng)
            each_w = (_emu(W, 0.46) - gap) // 2
            tl_f = self._jf(0.54, 0.02, rng)
            tw_f = self._jf(0.42, 0.02, rng)

        tl, tw = _emu(W, tl_f), _emu(W, tw_f)
        area_l = _emu(W, area_l_f)

        date_top = _emu(H, self._jf(0.04, 0.01, rng))
        title_top = _emu(H, self._jf(0.12, 0.01, rng))

        self._add_date(slide, summary.date, tl, date_top, tw)
        self._add_title_at(slide, summary.title, tl, title_top, tw)

        img_top = _emu(H, self._jf(0.10, 0.02, rng))
        img_h = _emu(H, self._jf(0.68, 0.03, rng))
        # Slight vertical stagger between the two images
        stagger = _emu(H, rng.uniform(0, 0.04))
        self._add_picture_fit(slide, images[0], area_l, img_top, each_w, img_h)
        self._add_picture_fit(slide, images[1], area_l + each_w + gap, img_top + stagger, each_w, img_h - stagger)

        narr_top = _emu(H, self._jf(0.26, 0.02, rng))
        end_y = self._add_narrative(slide, summary, tl, narr_top, tw, _emu(H, 0.28), rng)

        km_top = end_y + _emu(H, 0.02)
        km_h = _emu(H, 0.70) - km_top + _emu(H, 0.04)
        if km_h > 0:
            self._add_key_moments(slide, summary.key_moments, tl, km_top, tw, km_h, rng)

        self._add_quote(slide, summary.quote, tl, _emu(H, self._jf(0.76, 0.02, rng)), tw)

    # ------------------------------------------------------------------
    # Layout: gallery -- 3-4 images in 2x2 grid
    # ------------------------------------------------------------------

    def _layout_gallery(self, slide, summary: DailySummary,
                        images: list[Path], idx: int) -> None:
        W, H = self.scfg.W, self.scfg.H
        rng = random.Random(hash(summary.date) * 3 + 403)
        self._km_font_size = Pt(11)
        right_img = (idx % 2 == 0)

        if right_img:
            tl_f = self._jf(0.04, 0.02, rng)
            tw_f = self._jf(0.42, 0.02, rng)
            gl_f = self._jf(0.50, 0.02, rng)
            gw_f = 0.46
        else:
            gl_f = self._jf(0.04, 0.02, rng)
            gw_f = 0.46
            tl_f = self._jf(0.54, 0.02, rng)
            tw_f = self._jf(0.42, 0.02, rng)

        tl, tw = _emu(W, tl_f), _emu(W, tw_f)
        gl, gw = _emu(W, gl_f), _emu(W, gw_f)

        date_top = _emu(H, self._jf(0.04, 0.01, rng))
        title_top = _emu(H, self._jf(0.12, 0.01, rng))

        self._add_date(slide, summary.date, tl, date_top, tw)
        self._add_title_at(slide, summary.title, tl, title_top, tw)

        # Organic scattered image placement (not a grid)
        n = len(images)
        grid_top = _emu(H, self._jf(0.06, 0.02, rng))
        grid_h = _emu(H, 0.84)
        if n == 3:
            # One large + two small stacked
            big_w = int(gw * self._jf(0.58, 0.04, rng))
            big_h = int(grid_h * self._jf(0.60, 0.04, rng))
            sm_w = gw - big_w - _emu(W, 0.01)
            sm_h = (grid_h - _emu(H, 0.02)) // 2
            # Randomize which side gets the big image
            if rng.random() > 0.5:
                big_x = gl + _emu(W, rng.uniform(-0.01, 0.01))
                sm_x = gl + big_w + _emu(W, 0.01) + _emu(W, rng.uniform(-0.005, 0.005))
            else:
                sm_x = gl + _emu(W, rng.uniform(-0.005, 0.005))
                big_x = gl + sm_w + _emu(W, 0.01) + _emu(W, rng.uniform(-0.01, 0.01))
            big_y = grid_top + _emu(H, rng.uniform(-0.01, 0.02))
            self._add_picture_fit(slide, images[0], big_x, big_y, big_w, big_h)
            for j in range(2):
                sy = grid_top + j * (sm_h + _emu(H, 0.02)) + _emu(H, rng.uniform(-0.01, 0.01))
                self._add_picture_fit(slide, images[1 + j], sm_x, sy, sm_w, sm_h)
        else:  # 4 images
            # Varied sizes: two bigger on top, two smaller on bottom (or vice versa)
            gap_x = _emu(W, rng.uniform(0.008, 0.02))
            gap_y = _emu(H, rng.uniform(0.01, 0.03))
            split = self._jf(0.55, 0.08, rng)  # top row gets ~55% height
            top_h = int(grid_h * split)
            bot_h = grid_h - top_h - gap_y
            w1 = int(gw * self._jf(0.52, 0.06, rng))
            w2 = gw - w1 - gap_x
            w3 = int(gw * self._jf(0.48, 0.06, rng))
            w4 = gw - w3 - gap_x
            self._add_picture_fit(slide, images[0],
                gl + _emu(W, rng.uniform(-0.005, 0.005)),
                grid_top + _emu(H, rng.uniform(-0.005, 0.01)),
                w1, top_h)
            self._add_picture_fit(slide, images[1],
                gl + w1 + gap_x + _emu(W, rng.uniform(-0.005, 0.005)),
                grid_top + _emu(H, rng.uniform(-0.01, 0.01)),
                w2, top_h)
            bot_y = grid_top + top_h + gap_y
            self._add_picture_fit(slide, images[2],
                gl + _emu(W, rng.uniform(-0.005, 0.005)),
                bot_y + _emu(H, rng.uniform(-0.005, 0.005)),
                w3, bot_h)
            self._add_picture_fit(slide, images[3],
                gl + w3 + gap_x + _emu(W, rng.uniform(-0.005, 0.005)),
                bot_y + _emu(H, rng.uniform(-0.005, 0.01)),
                w4, bot_h)

        narr_top = _emu(H, self._jf(0.26, 0.02, rng))
        end_y = self._add_narrative(slide, summary, tl, narr_top, tw, _emu(H, 0.28), rng)

        km_top = end_y + _emu(H, 0.02)
        km_h = _emu(H, 0.70) - km_top + _emu(H, 0.04)
        if km_h > 0:
            self._add_key_moments(slide, summary.key_moments, tl, km_top, tw, km_h, rng)

        self._add_quote(slide, summary.quote, tl, _emu(H, self._jf(0.76, 0.02, rng)), tw)


def run_render(state: PipelineState) -> None:
    paths = state.paths
    summary_files = sorted(paths.daily_summaries_dir.glob('*.json'))
    logger.info('Found %d summary files to render', len(summary_files))

    if not summary_files:
        logger.warning('No summary files found. Cannot render PPTX.')
        return

    summaries = []
    for sf in summary_files:
        try:
            summaries.append(DailySummary.model_validate_json(sf.read_text(encoding='utf-8')))
        except Exception as exc:
            logger.error('Failed to load summary %s: %s', sf, exc)

    renderer = PptxRenderer(state)
    prs = renderer.render(summaries)

    paths.output_dir.mkdir(parents=True, exist_ok=True)
    prs.save(str(paths.pptx_output))
    logger.info('PPTX yearbook saved to: %s', paths.pptx_output)
